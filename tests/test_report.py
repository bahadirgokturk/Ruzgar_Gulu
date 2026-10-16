"""Faz 7: analiz zinciri ve Word rapor testleri."""
import re

import numpy as np
import pandas as pd
import pytest
from docx import Document

from core import economics as ec
from core import energy as en
from core.analysis import AnalysisInputs, run_analysis
from core.data_sources import clean_wind_data
from core.report import SECTION_TITLES, ReportMeta, build_report


@pytest.fixture(scope="module")
def result():
    rng = np.random.default_rng(3)
    n = 8760
    idx = pd.date_range("2024-01-01", periods=n, freq="h", tz="UTC")
    ws50 = 7.0 * rng.weibull(2.0, n)
    df = pd.DataFrame({"ws10": ws50 * 0.75, "ws50": ws50, "wd10": rng.uniform(0, 360, n), "wd50": rng.uniform(0, 360, n),
                       "t2m_c": 15 + 5 * rng.standard_normal(n), "ps_kpa": np.full(n, 100.5)}, index=idx)
    df, q = clean_wind_data(df)
    inp = AnalysisInputs(lat=38.42, lon=27.14, df=df, quality=q.as_dict(), source_label="sentetik veri",
                         turbine=en.load_library_turbine("E-82/2300", 98), econ=ec.EconomicInputs.from_defaults(),
                         env_notes={"noise": "En yakın yerleşim 1.2 km."})
    return run_analysis(inp)


@pytest.fixture(scope="module")
def doc_path(result, tmp_path_factory):
    meta = ReportMeta.from_defaults("Test Sahası", author="Deneme", include_map=False)
    return build_report(result, meta, out_dir=tmp_path_factory.mktemp("out"))


def test_analysis_consistency(result):
    assert result.econ is not None
    assert result.energy["net_aep"] > 0
    assert result.alpha_used == pytest.approx(np.log(1 / 0.75) / np.log(5), rel=1e-6)


def test_sections_present(doc_path):
    heads = [p.text for p in Document(doc_path).paragraphs if p.style.name == "Heading 1"]
    assert heads == [f"{i}. {t}" for i, t in enumerate(SECTION_TITLES, 1)]


def test_caption_numbering(doc_path):
    caps = [p.text for p in Document(doc_path).paragraphs if p.style.name == "Caption"]
    for kind in ("Şekil", "Tablo"):
        nums = [int(m.group(1)) for c in caps if (m := re.match(rf"{kind} (\d+)\.", c))]
        assert nums == list(range(1, len(nums) + 1)) and len(nums) >= 5


def test_summary_npv_matches(result, doc_path):
    text = "\n".join(p.text for p in Document(doc_path).paragraphs)
    assert f"{result.econ['npv'] / 1e6:.2f} milyon" in text
    assert "En yakın yerleşim 1.2 km." in text
    assert "değerlendirilmemiştir" in text            # boş nitel başlıklar


def test_cover_skips_empty_fields(doc_path):
    text = "\n".join(p.text for p in Document(doc_path).paragraphs[:30])
    assert "Hazırlayan: Deneme" in text
    assert "Öğretim Üyesi" not in text
