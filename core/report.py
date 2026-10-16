"""Word (.docx) fizibilite raporu üretimi, opsiyonel PDF dönüşümü.

Grafikler core/plots.py (matplotlib) ile PNG olarak üretilip gömülür.
Şekil ve tablo numaraları Python sayaçla verilir (Word alan güncellemesi gerektirmez).
"""
from __future__ import annotations

import datetime as dt
import io
import re
import shutil
import subprocess
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import matplotlib.pyplot as plt
import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from core import plots
from core.analysis import AnalysisResult
from core.config import load_defaults

OUTPUT_DIR = Path(__file__).resolve().parent.parent / "outputs"
MONTHS = ["Ocak", "Şubat", "Mart", "Nisan", "Mayıs", "Haziran", "Temmuz", "Ağustos", "Eylül", "Ekim", "Kasım", "Aralık"]
SECTION_TITLES = [
    "Yönetici Özeti", "Giriş ve Amaç", "Saha Tanımı", "Rüzgar Kaynak Değerlendirmesi",
    "Türbin Seçimi ve Enerji Üretim Tahmini", "Ekonomik Analiz", "Çevresel Etki",
    "Riskler ve Belirsizlikler", "Sonuç ve Öneriler", "Ekler",
]
REFERENCES = [
    "NASA Langley Research Center, POWER Project (Prediction Of Worldwide Energy Resources), https://power.larc.nasa.gov",
    "Haas, S. vd. windpowerlib: A Python library to model wind power plants, Journal of Open Research Software, 2024. "
    "https://github.com/wind-python/windpowerlib",
    "IEC 61400-12-1: Wind energy generation systems, Part 12-1: Power performance measurements of electricity "
    "producing wind turbines.",
    "Manwell, J. F., McGowan, J. G., Rogers, A. L. Wind Energy Explained: Theory, Design and Application, 2. baskı, Wiley, 2009.",
    "Elliott, D. L. vd. Wind Energy Resource Atlas of the United States, NREL/PNL, 1986.",
    "Justus, C. G. vd. Methods for estimating wind speed frequency distributions, Journal of Applied Meteorology, 1978.",
    "IRENA, Renewable Power Generation Costs in 2023, Abu Dhabi, 2024.",
    "T.C. Enerji ve Tabii Kaynaklar Bakanlığı, Elektrik Üretim ve Tüketim Noktası Emisyon Faktörleri, 2022.",
]
FORMULAS = [
    ("Hava yoğunluğu", "ρ = P / (R·T),  R = 287.05 J/(kg·K), P [Pa], T [K]"),
    ("Weibull olasılık yoğunluğu", "f(v) = (k/c)·(v/c)^(k−1)·exp(−(v/c)^k)"),
    ("Weibull, momentler yöntemi", "k = (σ/v̄)^(−1.086),  c = v̄ / Γ(1 + 1/k)"),
    ("Power law", "v₂ = v₁·(h₂/h₁)^α,  α = ln(v̄₅₀/v̄₁₀) / ln(50/10)"),
    ("Log law", "v₂ = v₁·ln(h₂/z₀) / ln(h₁/z₀)"),
    ("Güç yoğunluğu (zaman serisi)", "P/A = (1/N)·Σ ½·ρ·vₜ³"),
    ("Güç yoğunluğu (Weibull)", "P/A = ½·ρ·c³·Γ(1 + 3/k)"),
    ("Yoğunluk düzeltmesi (IEC 61400-12-1)", "v_saha = v_std·(ρ₀/ρ_saha)^(1/3),  ρ₀ = 1.225 kg/m³"),
    ("Brüt AEP (zaman serisi)", "AEP = 8760·mean(P(vₜ))"),
    ("Brüt AEP (Weibull)", "AEP = 8760·∫ P(v)·f(v) dv"),
    ("Net AEP", "Net = Brüt·Π(1 − kᵢ)"),
    ("Kapasite faktörü", "CF = Net AEP / (Pₙ·8760)"),
    ("LCOE", "LCOE = Σ (CAPEXₜ + OPEXₜ)/(1+r)ᵗ ÷ Σ Eₜ/(1+r)ᵗ"),
    ("NPV", "NPV = Σ NAₜ/(1+r)ᵗ,  t = 0..N"),
    ("IRR", "NPV(IRR) = 0"),
    ("AEP senaryoları", "P90 = P50·(1 − 1.282σ),  P10 = P50·(1 + 1.282σ)"),
    ("Önlenen CO₂", "CO₂ = Net AEP·EF"),
]


@dataclass
class ReportMeta:
    """Kapak ve genel bilgiler. Boş alanlar kapakta görünmez."""

    project: str
    author: str = ""
    university: str = ""
    department: str = ""
    course: str = ""
    instructor: str = ""
    student_no: str = ""
    date: dt.date = field(default_factory=dt.date.today)
    include_map: bool = True

    @classmethod
    def from_defaults(cls, project: str, **overrides: Any) -> "ReportMeta":
        cfg = load_defaults()["report"]
        vals = {k: cfg[k]["value"] for k in ("university", "department", "course", "instructor", "student_no",
                                              "include_map") if k in cfg}
        vals.update({k: v for k, v in overrides.items() if v is not None})
        return cls(project=project, **vals)


def fmt(x: float | None, d: int = 0) -> str:
    """Binlik ayırıcı boşluk, ondalık nokta."""
    if x is None or (isinstance(x, float) and pd.isna(x)):
        return "-"
    return f"{x:,.{d}f}".replace(",", " ")


def pct(x: float | None, d: int = 1) -> str:
    return "-" if x is None or pd.isna(x) else f"%{100 * x:.{d}f}"


# ---------------------------------------------------------------- belge yardımcıları

class _Doc:
    """Numaralı şekil/tablo ekleyen ince python-docx sarmalayıcısı."""

    def __init__(self) -> None:
        self.doc = Document()
        self.fig_n = 0
        self.tab_n = 0
        self._setup()

    def _setup(self) -> None:
        st = self.doc.styles
        st["Normal"].font.name = "Calibri"
        st["Normal"].font.size = Pt(11)
        st["Normal"].paragraph_format.space_after = Pt(6)
        for name in ("Heading 1", "Heading 2"):
            st[name].font.name = "Calibri"
            st[name].font.color.rgb = RGBColor(0x0E, 0x6B, 0x5A)
        st["Heading 1"].font.size = Pt(15)
        st["Heading 2"].font.size = Pt(12)
        sec = self.doc.sections[0]
        sec.left_margin = sec.right_margin = Cm(2.5)
        sec.top_margin = sec.bottom_margin = Cm(2.2)
        sec.different_first_page_header_footer = True
        p = sec.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
            if tag:
                el = OxmlElement("w:fldChar")
                el.set(qn("w:fldCharType"), tag)
            else:
                el = OxmlElement("w:instrText")
                el.set(qn("xml:space"), "preserve")
                el.text = text
            run._r.append(el)

    def h1(self, n: int, text: str) -> None:
        self.doc.add_heading(f"{n}. {text}", level=1)

    def h2(self, text: str) -> None:
        self.doc.add_heading(text, level=2)

    def p(self, text: str, bold: bool = False, italic: bool = False, align=None) -> None:
        para = self.doc.add_paragraph()
        run = para.add_run(text)
        run.bold, run.italic = bold, italic
        if align is not None:
            para.alignment = align

    def bullets(self, items: list[str]) -> None:
        for it in items:
            self.doc.add_paragraph(it, style="List Bullet")

    def _caption(self, text: str) -> None:
        para = self.doc.add_paragraph(style="Caption")
        para.add_run(text)
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def figure(self, fig_or_png, caption: str, width_cm: float = 15) -> int:
        """Şekil ekler, numarasını döndürür. fig_or_png: matplotlib Figure veya PNG baytları."""
        if isinstance(fig_or_png, (bytes, bytearray)):
            buf = io.BytesIO(fig_or_png)
        else:
            buf = io.BytesIO()
            fig_or_png.savefig(buf, format="png", dpi=150, bbox_inches="tight")
            plt.close(fig_or_png)
            buf.seek(0)
        self.doc.add_picture(buf, width=Cm(width_cm))
        self.doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        self.fig_n += 1
        self._caption(f"Şekil {self.fig_n}. {caption}")
        return self.fig_n

    def table(self, rows: list[list[str]], caption: str, header: list[str] | None = None,
              widths_cm: list[float] | None = None) -> int:
        """Tablo ekler (başlık tablonun üstünde), numarasını döndürür."""
        self.tab_n += 1
        self._caption(f"Tablo {self.tab_n}. {caption}")
        data = ([header] if header else []) + rows
        t = self.doc.add_table(rows=len(data), cols=len(data[0]))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, row in enumerate(data):
            for j, val in enumerate(row):
                cell = t.cell(i, j)
                cell.text = str(val)
                for para in cell.paragraphs:
                    para.paragraph_format.space_after = Pt(0)
                    for r in para.runs:
                        r.font.size = Pt(9.5)
                        r.bold = bool(header) and i == 0
                if header and i == 0:
                    shd = OxmlElement("w:shd")
                    shd.set(qn("w:val"), "clear")
                    shd.set(qn("w:fill"), "E2EEEA")
                    cell._tc.get_or_add_tcPr().append(shd)
        if widths_cm:
            for row in t.rows:
                for cell, w in zip(row.cells, widths_cm):
                    cell.width = Cm(w)
        self.doc.add_paragraph()
        return self.tab_n

    def page_break(self) -> None:
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def site_map_png(lat: float, lon: float, zoom: int = 9) -> bytes | None:
    """OpenStreetMap karolarından konum işaretli statik harita (PNG). İnternet yoksa None."""
    try:
        from staticmap import CircleMarker, StaticMap

        m = StaticMap(900, 520, url_template="https://tile.openstreetmap.org/{z}/{x}/{y}.png",
                      headers={"User-Agent": "ruzgar-fizibilite-ders-projesi/1.0"})
        m.add_marker(CircleMarker((lon, lat), "white", 18))
        m.add_marker(CircleMarker((lon, lat), "#0e6b5a", 12))
        img = m.render(zoom=zoom)
        buf = io.BytesIO()
        img.save(buf, format="PNG")
        return buf.getvalue()
    except Exception:
        return None


# ---------------------------------------------------------------- otomatik metinler

def summary_sentences(r: AnalysisResult) -> list[str]:
    """Yönetici özeti için otomatik cümleler."""
    from core.economics import verdict

    i, e = r.inputs, r.energy
    out = [
        f"{i.lat:.4f} K, {i.lon:.4f} D konumundaki saha için {i.source_label} verisiyle yapılan analizde, "
        f"{i.turbine.hub_height:g} m göbek yüksekliğinde ortalama rüzgar hızı {r.ws_hub.mean():.2f} m/s, "
        f"güç yoğunluğu {fmt(r.pd_hub)} W/m² olarak hesaplanmıştır. Saha, {i.h1:g} m'de NREL sınıflamasına göre "
        f"{r.wind_class['class']}. sınıf ({r.wind_class['name'].lower()}) rüzgar kaynağına sahiptir.",
        f"{i.n_turbines} adet {i.turbine.name} ({fmt(i.turbine.rated_kw)} kW) türbin ile toplam "
        f"{e['capacity_mw']:.2f} MW kurulu güçte yıllık net üretim {fmt(e['net_aep'])} MWh, kapasite faktörü "
        f"{pct(e['capacity_factor'])} ve eşdeğer tam yük saati {fmt(e['full_load_hours'])} h olarak tahmin edilmiştir.",
    ]
    if r.econ:
        c = r.econ
        pb = "proje ömrünü aşan bir süre" if c["payback"] is None else f"{c['payback']:.1f} yıl"
        out.append(
            f"Seviyelendirilmiş enerji maliyeti (LCOE) {c['lcoe']:.1f} {i.econ.currency}/MWh, %{100 * i.econ.discount_rate:.0f} "
            f"iskonto oranıyla net bugünkü değer {c['npv'] / 1e6:.2f} milyon {i.econ.currency}, iç verim oranı "
            f"{pct(c['irr'])} ve basit geri ödeme süresi {pb} olarak hesaplanmıştır.")
        out.append(verdict(c, i.econ.discount_rate)[1])
    out.append(f"Proje yılda yaklaşık {fmt(r.env['annual_co2_t'])} ton CO₂ salımını önleyecek ve "
               f"yaklaşık {fmt(r.env['households'])} hanenin yıllık elektrik tüketimini karşılayacaktır.")
    return out


def risk_items(r: AnalysisResult) -> list[str]:
    """Sonuçlara göre koşullu risk ve belirsizlik maddeleri."""
    i = r.inputs
    items = []
    if i.is_nasa:
        items.append("Rüzgar verisi NASA POWER reanaliz veri setinden alınmıştır. Reanaliz verisi yerel arazi ve "
                     "engel etkilerini yumuşattığı için sahadaki gerçek hızdan sapabilir. Bankaya sunulabilir bir "
                     "değerlendirme için en az bir yıllık yerinde ölçüm gereklidir.")
    else:
        items.append("Rüzgar verisi kullanıcının ölçümüdür. Ölçüm süresi bir yıldan kısaysa uzun dönem "
                     "korelasyonu (MCP) yapılmadan sonuçlar mevsimsel olarak yanlı olabilir.")
    if r.inputs.quality.get("missing_pct", 0) > 5:
        items.append(f"Veride %{i.quality['missing_pct']:.1f} eksik kayıt vardır; bu oran yıllık tahmini etkileyebilir.")
    if r.method_compare and abs(r.method_compare["diff_pct"]) > 2:
        items.append(f"Göbek yüksekliğine ekstrapolasyonda power law ve log law arasında "
                     f"%{abs(r.method_compare['diff_pct']):.1f} fark vardır. Rüzgar profili (α, z₀) varsayımı "
                     f"üretim tahminini doğrudan etkiler.")
    if r.rho_is_standard:
        items.append("Veride basınç ve sıcaklık bulunmadığından standart hava yoğunluğu (1.225 kg/m³) kullanılmıştır.")
    else:
        items.append("Hava yoğunluğu yüzey basıncı ve 2 m sıcaklıkla hesaplanmış, göbek yüksekliğine düzeltilmemiştir.")
    items.append(f"Kayıp oranları (toplam %{r.energy['total_loss_pct']:.1f}) sahaya özel değil, tipik değerlerdir.")
    if r.sensitivity is not None:
        s = r.sensitivity
        rng = s.groupby("label")["d_npv"].agg(lambda x: x.max() - x.min()).sort_values(ascending=False)
        items.append(f"Duyarlılık analizine göre NPV'yi en çok etkileyen parametreler {rng.index[0]} ve "
                     f"{rng.index[1]}'dır. Bu iki büyüklükteki belirsizlik sonucu belirler.")
    items.append("Satış fiyatı ve teşvik dönemi kullanıcı varsayımıdır; güncel mevzuat ve piyasa koşullarına göre "
                 "değişebilir.")
    return items


def recommendations(r: AnalysisResult) -> list[str]:
    """Fizibilite sonucuna göre öneriler."""
    rec = []
    ok = bool(r.econ and r.econ["npv"] > 0)
    if ok:
        rec.append("Ön fizibilite olumludur. Bir sonraki adım olarak sahaya en az 12 ay sürecek bir ölçüm direği "
                   "kurulması ve sonuçların yerinde ölçümle doğrulanması önerilir.")
    else:
        rec.append("Mevcut varsayımlarla proje ekonomik olarak fizibil görünmemektedir. Karar vermeden önce rüzgar "
                   "kaynağının yerinde ölçümle doğrulanması önerilir; reanaliz verisi kaynağı düşük gösteriyor olabilir.")
        rec.append("Daha yüksek göbek yüksekliği, daha büyük rotor çaplı (düşük rüzgar sınıfı) bir türbin veya daha "
                   "fazla türbinle ölçek ekonomisi denenebilir.")
    if r.energy["capacity_factor"] < 0.25:
        rec.append("Kapasite faktörü %25'in altındadır; seçilen türbinin saha rüzgar rejimine uygunluğu "
                   "yeniden değerlendirilmelidir.")
    rec.append("Çevresel etkiler (gürültü, gölge titreşimi, kuş göçü) için ÇED sürecinde ayrıntılı çalışma yapılmalıdır.")
    return rec


# ---------------------------------------------------------------- rapor

def build_report(r: AnalysisResult, meta: ReportMeta, out_dir: Path = OUTPUT_DIR,
                 map_png: bytes | None = None) -> Path:
    """Türkçe fizibilite raporunu .docx olarak üretir ve yolunu döndürür."""
    i, e, D = r.inputs, r.energy, _Doc()
    cur = i.econ.currency
    doc = D.doc

    # ---- kapak
    for _ in range(4):
        doc.add_paragraph()
    for line in (meta.university, meta.department, meta.course and f"{meta.course} Dersi"):
        if line:
            D.p(line, align=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_paragraph()
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = t.add_run("RÜZGAR ENERJİSİ SANTRALİ\nFİZİBİLİTE RAPORU")
    run.bold, run.font.size = True, Pt(20)
    run.font.color.rgb = RGBColor(0x0E, 0x6B, 0x5A)
    D.p(meta.project, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    D.p(f"{i.lat:.4f} K, {i.lon:.4f} D", align=WD_ALIGN_PARAGRAPH.CENTER)
    for _ in range(5):
        doc.add_paragraph()
    for label, val in (("Hazırlayan", meta.author), ("Öğrenci No", meta.student_no),
                       ("Öğretim Üyesi", meta.instructor), ("Tarih", meta.date.strftime("%d.%m.%Y"))):
        if val:
            D.p(f"{label}: {val}", align=WD_ALIGN_PARAGRAPH.CENTER)
    D.page_break()

    # ---- 1. yönetici özeti
    D.h1(1, SECTION_TITLES[0])
    for s in summary_sentences(r):
        D.p(s)
    rows = [
        [f"Ortalama hız ({i.turbine.hub_height:g} m)", f"{r.ws_hub.mean():.2f} m/s"],
        ["Güç yoğunluğu (göbek)", f"{fmt(r.pd_hub)} W/m²"],
        ["Kurulu güç", f"{e['capacity_mw']:.2f} MW"],
        ["Net AEP", f"{fmt(e['net_aep'])} MWh/yıl"],
        ["Kapasite faktörü", pct(e["capacity_factor"])],
    ]
    if r.econ:
        rows += [["LCOE", f"{r.econ['lcoe']:.1f} {cur}/MWh"], ["NPV", f"{r.econ['npv'] / 1e6:.2f} M {cur}"],
                 ["IRR", pct(r.econ["irr"])],
                 ["Geri ödeme (basit)", "-" if r.econ["payback"] is None else f"{r.econ['payback']:.1f} yıl"]]
    rows.append(["Önlenen CO₂", f"{fmt(r.env['annual_co2_t'])} t/yıl"])
    D.table(rows, "Ana sonuçlar", ["Gösterge", "Değer"], [8, 6])

    # ---- 2. giriş
    D.h1(2, SECTION_TITLES[1])
    D.p("Bu rapor, seçilen sahada kurulması düşünülen rüzgar enerjisi santralinin teknik ve ekonomik ön "
        "fizibilitesini değerlendirmek amacıyla hazırlanmıştır. Çalışmada saatlik rüzgar verisi istatistiksel olarak "
        "incelenmiş, Weibull dağılımı ile modellenmiş, göbek yüksekliğine taşınmış ve seçilen türbinin güç eğrisiyle "
        "yıllık enerji üretimi tahmin edilmiştir. Elde edilen üretim; yatırım, işletme maliyetleri ve satış fiyatı "
        "varsayımlarıyla birlikte nakit akışı analizine aktarılarak LCOE, NPV ve IRR göstergeleri hesaplanmıştır.")
    D.p("Hesaplamalar, bu ders kapsamında geliştirilen açık kaynaklı Python yazılımı ile yapılmıştır. Tüm varsayımlar "
        "ve kaynakları Ekler bölümünde verilmiştir.")

    # ---- 3. saha
    D.h1(3, SECTION_TITLES[2])
    D.table([["Enlem", f"{i.lat:.4f}°"], ["Boylam", f"{i.lon:.4f}°"],
             ["Arazi tipi", i.z0_name], ["Pürüzlülük uzunluğu z₀", f"{i.z0:g} m"]],
            "Saha bilgileri", ["Özellik", "Değer"], [8, 6])
    png = map_png if map_png is not None else (site_map_png(i.lat, i.lon) if meta.include_map else None)
    if png:
        D.figure(png, "Saha konumu (harita: © OpenStreetMap katkıcıları)")

    # ---- 4. rüzgar kaynağı
    D.h1(4, SECTION_TITLES[3])
    D.h2("4.1 Veri kaynağı ve kalitesi")
    q = i.quality
    D.p(f"Analizde {i.source_label} kullanılmıştır. Veri temizleme adımında tekrar eden zaman damgaları, negatif ve "
        f"50 m/s üzerindeki hızlar ile 0-360° dışındaki yönler ayıklanmıştır.")
    D.table([["Kapsanan süre", f"{pd.Timestamp(q['start']):%d.%m.%Y} - {pd.Timestamp(q['end']):%d.%m.%Y}"],
             ["Toplam kayıt", fmt(q["total_records"])], ["Beklenen kayıt", fmt(q["expected_records"])],
             ["Eksik veri", f"%{q['missing_pct']:.2f}"], ["Tekrarlı zaman damgası", fmt(q["duplicates_removed"])],
             ["Fiziksel olmayan değer", fmt(q["out_of_range_removed"])]],
            "Veri kalite özeti", ["Ölçüt", "Değer"], [8, 6])
    D.h2("4.2 Temel istatistikler ve Weibull dağılımı")
    s = r.stats
    D.table([["Ortalama", f"{s['mean']:.2f} m/s"], ["Medyan", f"{s['median']:.2f} m/s"],
             ["Standart sapma", f"{s['std']:.2f} m/s"], ["En yüksek", f"{s['max']:.2f} m/s"],
             ["Sakin saatler (< 0.5 m/s)", f"%{s['calm_pct']:.1f}"],
             ["Yıllar arası değişkenlik (CoV)", f"%{r.cov:.1f}"], ["Ortalama hava yoğunluğu", f"{r.rho:.3f} kg/m³"]],
            f"{i.h1:g} m rüzgar hızı istatistikleri", ["Büyüklük", "Değer"], [8, 6])
    m, mo = r.weibull_mle, r.weibull_mom
    from core.wind_resource import power_density_weibull
    D.table([["Maximum Likelihood", f"{m.k:.2f}", f"{m.c:.2f}", fmt(power_density_weibull(m.k, m.c, r.rho))],
             ["Momentler (ampirik)", f"{mo.k:.2f}", f"{mo.c:.2f}", fmt(power_density_weibull(mo.k, mo.c, r.rho))],
             ["Zaman serisi, ½ρv³", "-", "-", fmt(r.pd_ref)]],
            "Weibull parametreleri ve güç yoğunluğu", ["Yöntem", "k", "c (m/s)", "P/A (W/m²)"], [6, 2.5, 2.5, 3])
    n = D.figure(plots.weibull_histogram(i.df[i.ws_col]), f"{i.h1:g} m hız histogramı ve Weibull uyumu")
    D.p(f"Şekil {n}'de görüldüğü gibi iki yöntem birbirine yakın sonuç vermektedir. Güç yoğunluğu ortalama hızın "
        f"küpüyle değil, saatlik ½ρv³ değerlerinin ortalaması olarak hesaplanmıştır. Saha {i.h1:g} m'de "
        f"{fmt(r.pd_ref)} W/m² ile NREL sınıflamasında {r.wind_class['class']}. sınıftadır "
        f"({r.wind_class['name'].lower()}).")
    if r.rose is not None:
        top = r.rose.sort_values("freq_pct").iloc[-1]
        n = D.figure(plots.wind_rose_plot(r.rose), "Rüzgar gülü (16 sektör, frekans)", width_cm=10)
        D.p(f"Şekil {n}'deki rüzgar gülüne göre baskın yön {top['sector']} (%{top['freq_pct']:.1f}) olup bu "
            f"sektörde ortalama hız {top['mean_ws']:.2f} m/s'dir.")
    mp = r.monthly
    D.table([[MONTHS[k - 1], f"{v:.2f}"] for k, v in mp.items()], "Aylık ortalama rüzgar hızı",
            ["Ay", "Ortalama hız (m/s)"], [6, 6])
    D.h2("4.3 Göbek yüksekliğine ekstrapolasyon")
    method_txt = (f"power law (α = {r.alpha_used:.3f})" if i.method == "power" else f"log law (z₀ = {i.z0:g} m)")
    D.p(f"{i.h1:g} m verisi {i.turbine.hub_height:g} m göbek yüksekliğine {method_txt} ile taşınmıştır.")
    if r.method_compare:
        c = r.method_compare
        D.table([[f"Power law (α = {c['alpha']:.3f})", f"{c['power_mean']:.2f}"],
                 [f"Log law (z₀ = {c['z0']:g} m)", f"{c['log_mean']:.2f}"]],
                f"{i.turbine.hub_height:g} m ortalama hız, yöntem karşılaştırması", ["Yöntem", "Ortalama hız (m/s)"], [8, 6])
        D.p(f"İki yöntem arasındaki fark %{abs(c['diff_pct']):.1f}'dir.")

    # ---- 5. türbin ve enerji
    D.h1(5, SECTION_TITLES[4])
    t_ = i.turbine
    D.table([["Model", t_.name], ["Nominal güç", f"{fmt(t_.rated_kw)} kW"],
             ["Rotor çapı", "-" if t_.rotor_d is None else f"{t_.rotor_d:g} m"],
             ["Göbek yüksekliği", f"{t_.hub_height:g} m"], ["Türbin adedi", str(i.n_turbines)],
             ["Yoğunluk düzeltmesi", "IEC 61400-12-1" if i.density_method == "iec" else "Svenningsen"]],
            "Türbin bilgileri", ["Özellik", "Değer"], [8, 6])
    n = D.figure(plots.power_curve_plot(e["curve_std"], e["curve_site"]), "Güç eğrisi, standart ve saha yoğunluğu")
    D.p(f"Şekil {n}'de güç eğrisi saha ortalama hava yoğunluğuna ({r.rho:.3f} kg/m³) göre düzeltilmiştir.")
    D.table([["Zaman serisi", fmt(e["gross_aep_ts"])], ["Weibull integrali", fmt(e["gross_aep_weibull"])]],
            "Brüt AEP, yöntem karşılaştırması", ["Yöntem", "Brüt AEP (MWh/yıl)"], [8, 6])
    D.p(f"İki yöntem arasındaki fark %{abs(e['method_diff_pct']):.1f}'dir. Net üretim zaman serisi yöntemine göre hesaplanmıştır.")
    lt = e["loss_table"]
    D.table([[row.label, f"%{100 * row.rate:.1f}", fmt(row.lost_mwh)] for row in lt.itertuples()]
            + [["Toplam (çarpımsal)", f"%{e['total_loss_pct']:.1f}", fmt(e["gross_aep_ts"] - e["net_aep"])]],
            "Enerji kayıpları", ["Kalem", "Oran", "Kayıp (MWh/yıl)"], [7, 3, 4])
    D.figure(plots.loss_waterfall(e["gross_aep_ts"], lt[lt["rate"] > 0], e["net_aep"]), "Brüttan nete kayıp şelalesi")
    D.table([["Net AEP", f"{fmt(e['net_aep'])} MWh/yıl"], ["Kapasite faktörü", pct(e["capacity_factor"])],
             ["Eşdeğer tam yük saati", f"{fmt(e['full_load_hours'])} h"]],
            "Enerji üretim özeti", ["Gösterge", "Değer"], [8, 6])
    D.table([[MONTHS[k - 1], fmt(v)] for k, v in e["monthly_net"].items()], "Aylık ortalama net üretim",
            ["Ay", "Net üretim (MWh)"], [6, 6])

    # ---- 6. ekonomi
    D.h1(6, SECTION_TITLES[5])
    x = i.econ
    D.table([["CAPEX", f"{fmt(x.capex_per_kw)} {cur}/kW"], ["OPEX", f"{fmt(x.opex_per_kw_year)} {cur}/kW/yıl"],
             ["OPEX artışı", pct(x.opex_escalation)],
             ["Teşvik fiyatı / süresi", f"{fmt(x.incentive_price, 1)} {cur}/MWh, {x.incentive_years} yıl"],
             ["Teşvik sonrası fiyat", f"{fmt(x.market_price, 1)} {cur}/MWh (eskalasyon {pct(x.price_escalation)})"],
             ["İskonto oranı", pct(x.discount_rate)], ["Proje ömrü", f"{x.lifetime} yıl"],
             ["Degradasyon", f"{pct(x.degradation)}/yıl"], ["Kurumlar vergisi", pct(x.tax_rate)],
             ["Kredi", f"%{100 * x.loan_share:.0f} pay, {pct(x.loan_rate)} faiz, {x.loan_term} yıl" if x.loan_enabled else "Yok"]],
            "Ekonomik varsayımlar", ["Varsayım", "Değer"], [7, 7])
    if r.econ is None:
        D.p(f"Ekonomik analiz yapılamadı: {r.econ_error}")
    else:
        c = r.econ
        D.table([[k, fmt(v / 1e3)] for k, v in c["capex_breakdown"].items()] + [["Toplam", fmt(c["capex"] / 1e3)]],
                "Yatırım maliyeti dağılımı", ["Kalem", f"Tutar (bin {cur})"], [8, 6])
        T = c["table"]
        D.table([[str(y), fmt(T.loc[y, "energy_mwh"]), fmt(T.loc[y, "revenue"] / 1e3), fmt(T.loc[y, "opex"] / 1e3),
                  fmt(T.loc[y, "tax"] / 1e3), fmt(T.loc[y, "net_cf"] / 1e3), fmt(T.loc[y, "cum_cf"] / 1e3)]
                 for y in T.index],
                f"Yıllık nakit akışı (bin {cur})", ["Yıl", "Üretim (MWh)", "Gelir", "OPEX", "Vergi", "Net NA", "Kümülatif"])
        D.table([["LCOE", f"{c['lcoe']:.1f} {cur}/MWh"], ["NPV", f"{fmt(c['npv'])} {cur}"], ["IRR", pct(c["irr"])],
                 ["Basit geri ödeme", "-" if c["payback"] is None else f"{c['payback']:.1f} yıl"],
                 ["İskontolu geri ödeme", "ömür içinde yok" if c["discounted_payback"] is None
                  else f"{c['discounted_payback']:.1f} yıl"]],
                "Ekonomik göstergeler", ["Gösterge", "Değer"], [8, 6])
        D.figure(plots.cash_flow_plot(T, cur), "Yıllık ve kümülatif nakit akışı")
        n = D.figure(plots.tornado_plot(r.sensitivity, currency=cur), "NPV duyarlılık analizi (±%20)")
        D.p(f"Şekil {n}, parametrelerin tek tek ±%20 değiştirilmesinin NPV üzerindeki etkisini göstermektedir.")
        sc = r.scenarios
        D.table([[idx, row.p_level, fmt(row.net_aep), f"{row.npv / 1e6:.2f}", pct(row.irr), f"{row.lcoe:.1f}"]
                 for idx, row in sc.iterrows()],
                f"Senaryo analizi (σ = %{100 * x.aep_sigma:.0f})",
                ["Senaryo", "AEP", "Net AEP (MWh)", f"NPV (M {cur})", "IRR", f"LCOE ({cur}/MWh)"])

    # ---- 7. çevre
    D.h1(7, SECTION_TITLES[6])
    ev = r.env
    D.table([["Şebeke emisyon faktörü", f"{ev['emission_factor']:.3f} tCO₂/MWh"],
             ["Yıllık önlenen CO₂", f"{fmt(ev['annual_co2_t'])} t"],
             [f"{x.lifetime} yıllık toplam", f"{fmt(ev['lifetime_co2_t'])} t"],
             ["Karşılanan hane sayısı", f"≈ {fmt(ev['households'])} ({fmt(ev['household_kwh'])} kWh/yıl/hane)"]],
            "Çevresel kazanımlar", ["Gösterge", "Değer"], [7, 7])
    for title, text in ev["qualitative"].values():
        D.h2(title)
        D.p(text if text else "Bu başlık bu çalışma kapsamında değerlendirilmemiştir.", italic=not text)

    # ---- 8. riskler
    D.h1(8, SECTION_TITLES[7])
    D.bullets(risk_items(r))

    # ---- 9. sonuç
    D.h1(9, SECTION_TITLES[8])
    if r.econ:
        from core.economics import verdict
        D.p(verdict(r.econ, x.discount_rate)[1], bold=True)
    D.bullets(recommendations(r))

    # ---- 10. ekler
    D.h1(10, SECTION_TITLES[9])
    D.h2("Ek A. Varsayılan değerler")
    rows = []
    for sec, items in load_defaults().items():
        for k, v in (items or {}).items():
            if not isinstance(v, dict) or "value" not in v:
                continue
            val = v["value"]
            val = "tablo" if isinstance(val, dict) else ", ".join(map(str, val)) if isinstance(val, list) else str(val)
            rows.append([f"{sec}.{k}", f"{val} {v.get('unit', '')}".strip(), str(v.get("source", ""))])
    D.table(rows, "Varsayılan değerler ve kaynakları (config/defaults.yaml)", ["Parametre", "Değer", "Kaynak"], [5, 3.5, 7.5])
    D.h2("Ek B. Kullanılan formüller")
    D.table([[a, b] for a, b in FORMULAS], "Formüller", ["Büyüklük", "Formül"], [5, 11])
    D.h2("Ek C. Kaynakça")
    for k, ref in enumerate(REFERENCES, 1):
        D.p(f"[{k}] {ref}")

    out_dir.mkdir(parents=True, exist_ok=True)
    safe = re.sub(r"[^\w\-]+", "_", meta.project, flags=re.UNICODE).strip("_") or "rapor"
    path = out_dir / f"{safe}_{meta.date:%Y%m%d}.docx"
    doc.save(path)
    return path


def to_pdf(docx_path: Path) -> Path | None:
    """docx → PDF: önce docx2pdf (MS Word), olmazsa LibreOffice. İkisi de yoksa None."""
    pdf = docx_path.with_suffix(".pdf")
    try:
        from docx2pdf import convert

        convert(str(docx_path), str(pdf))
        if pdf.exists():
            return pdf
    except Exception:
        pass
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if soffice:
        try:
            subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(docx_path.parent),
                            str(docx_path)], check=True, capture_output=True, timeout=180)
            if pdf.exists():
                return pdf
        except Exception:
            pass
    return None
